"""
PDF → PPTX コンバーター v4.1
ライトモード + ハードモード（Gemini AI）
"""

import io, tempfile, os, base64, json, re
import streamlit as st

st.set_page_config(
    page_title="PDF → PPTX コンバーター",
    page_icon="📄",
    layout="centered",
    initial_sidebar_state="collapsed",
)

st.markdown("""
<style>
html, body, [data-testid="stAppViewContainer"] { background: #f0f2f6; }
[data-testid="stHeader"] { background: transparent; }
.main .block-container { padding-top: 1.5rem; padding-bottom: 3rem; max-width: 760px; }
.app-header {
    background: white; border-radius: 16px; padding: 1.5rem 2rem;
    margin-bottom: 1.5rem; box-shadow: 0 2px 12px rgba(0,0,0,0.06);
    border-left: 5px solid #667eea;
}
.app-header h1 { font-size: 1.4rem; font-weight: 700; color: #1a1a2e; margin: 0 0 0.3rem; }
.app-header p  { font-size: 0.85rem; color: #666; margin: 0; }
.section-card {
    background: white; border-radius: 14px; padding: 1.4rem 1.5rem;
    margin-bottom: 1rem; box-shadow: 0 2px 8px rgba(0,0,0,0.04);
}
.section-card h3 { font-size: 1rem; font-weight: 700; color: #333; margin: 0 0 1rem; }
.mode-info {
    background: #f5f3ff; border-radius: 10px; padding: 1rem 1.2rem;
    margin-top: 0.8rem; font-size: 0.83rem; color: #4a4a8a; line-height: 1.7;
    border: 1px solid #e0d9ff;
}
.api-info {
    background: #fff8e1; border-radius: 10px; padding: 0.8rem 1rem;
    margin: 0.8rem 0; font-size: 0.83rem; color: #e65100;
    border-left: 3px solid #ffa000;
}
.stButton > button {
    background: linear-gradient(135deg, #667eea 0%, #764ba2 100%) !important;
    color: white !important; border: none !important;
    border-radius: 12px !important; padding: 0.75rem 2rem !important;
    font-size: 1rem !important; font-weight: 600 !important; width: 100% !important;
    box-shadow: 0 4px 15px rgba(102,126,234,0.35) !important;
}
[data-testid="stDownloadButton"] > button {
    background: white !important; color: #667eea !important;
    border: 2px solid #667eea !important; border-radius: 12px !important;
    width: 100% !important; font-size: 1rem !important;
    font-weight: 600 !important; padding: 0.75rem 2rem !important;
}
[data-testid="stDownloadButton"] > button:hover { background: #667eea !important; color: white !important; }
.stat-grid { display: grid; grid-template-columns: repeat(3,1fr); gap: 10px; margin: 1rem 0; }
.stat-box { background: #f8f9ff; border-radius: 10px; padding: 1rem; text-align: center; border: 1px solid #e8eaf0; }
.stat-num { font-size: 1.8rem; font-weight: 700; color: #667eea; }
.stat-lbl { font-size: 0.72rem; color: #888; margin-top: 2px; }
.stProgress > div > div > div { background: linear-gradient(90deg,#667eea,#764ba2) !important; border-radius: 4px !important; }
hr { border-color: #e8eaf0 !important; }
</style>
""", unsafe_allow_html=True)

# ヘッダー
st.markdown("""
<div class="app-header">
  <h1>📄 PDF → PPTX コンバーター</h1>
  <p>PDFスライドを、テキスト・画像・図形が編集可能なPowerPointに変換します</p>
</div>
""", unsafe_allow_html=True)

# ── ステップ① モード選択 ──────────────────────────────
st.markdown('<div class="section-card"><h3>ステップ① 変換モードを選んでください</h3>', unsafe_allow_html=True)

mode = st.radio(
    "モード選択",
    options=["⚡ ライトモード", "🤖 ハードモード（AI）"],
    index=0,
    horizontal=True,
    label_visibility="collapsed",
)
is_hard = "ハード" in mode

if is_hard:
    st.markdown("""
<div class="mode-info">
<b>🤖 ハードモード</b>：Gemini AIが画像を解析してテキスト・画像・図形を分離します<br>
✅ 画像に焼き付いたテキストも認識・編集可能に<br>
✅ 写真・イラストは図形として切り出して配置<br>
✅ フォントサイズも近似値で再現<br>
⏱ 1ページあたり15〜30秒かかります
</div>
""", unsafe_allow_html=True)
else:
    st.markdown("""
<div class="mode-info">
<b>⚡ ライトモード</b>：PDFからテキスト・図形・画像を直接抽出します<br>
✅ テキストが埋め込まれたPDFに最適・数秒で完了<br>
✅ 図形・背景色・画像を忠実に再現<br>
❌ 画像として書き出されたスライド（NotebookLMなど）はハードモードを使用
</div>
""", unsafe_allow_html=True)

st.markdown('</div>', unsafe_allow_html=True)

# ── ステップ② APIキー（ハードのみ） ──────────────────
gemini_key = ""
if hasattr(st, "secrets"):
    gemini_key = st.secrets.get("GEMINI_API_KEY", "")

if is_hard:
    st.markdown('<div class="section-card"><h3>ステップ② APIキーを入力</h3>', unsafe_allow_html=True)
    st.markdown("""
<div class="api-info">
🔑 Google AI Studio（無料）のAPIキーが必要です<br>
取得方法：<a href="https://aistudio.google.com" target="_blank">aistudio.google.com</a> を開く
→ 左メニュー「Get API key」→「Create API key」→ コピーして貼り付け
</div>
""", unsafe_allow_html=True)
    if not gemini_key:
        gemini_key = st.text_input(
            "APIキー（AIzaSy...で始まる文字列）",
            type="password",
            placeholder="AIzaSy...",
        )
    else:
        st.success("✅ APIキーは設定済みです")
    st.markdown('</div>', unsafe_allow_html=True)

# ── ステップ③ ファイルアップロード ───────────────────
step_num = "③" if is_hard else "②"
st.markdown(f'<div class="section-card"><h3>ステップ{step_num} PDFをアップロード</h3>', unsafe_allow_html=True)

uploaded = st.file_uploader(
    "PDFをドロップ、またはクリックして選択",
    type=["pdf"],
    help="スライド形式のPDF推奨（最大200MB）",
)

def fmt_size(n):
    if n < 1024**2: return f"{n/1024:.1f} KB"
    return f"{n/1024**2:.1f} MB"

if uploaded:
    st.success(f"📄 **{uploaded.name}** ({fmt_size(uploaded.size)}) 読み込み完了")
st.markdown('</div>', unsafe_allow_html=True)


# ── 変換エンジン（ライト） ───────────────────────────
def convert_light(pdf_bytes, slide_width_inches=13.333, progress_cb=None):
    import fitz
    from pptx import Presentation
    from pptx.util import Pt, Inches
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from lxml import etree
    from collections import Counter
    from PIL import Image as PILImage

    PT_TO_EMU = 914400 / 72
    JP_FONT_MAP = {
        "kozminpr6n":"游明朝","kozgopr6n":"游ゴシック",
        "hirakakupro":"ヒラギノ角ゴ Pro W3","msgothic":"ＭＳ ゴシック",
        "mspgothic":"ＭＳ Ｐゴシック","msmincho":"ＭＳ 明朝",
        "yugothic":"游ゴシック","yumincho":"游明朝",
        "heiseikakugo":"ＭＳ ゴシック","meiryo":"メイリオ",
    }

    def _to_rgb(c):
        if c is None: return None
        if isinstance(c, float) and 0<=c<=1: g=int(c*255); return (g,g,g)
        if isinstance(c, int): return ((c>>16)&0xFF,(c>>8)&0xFF,c&0xFF)
        if isinstance(c,(tuple,list)):
            if len(c)==3: return tuple(max(0,min(255,int(x*255))) for x in c)
            if len(c)==4:
                C,M,Y,K=c
                return (int(255*(1-C)*(1-K)),int(255*(1-M)*(1-K)),int(255*(1-Y)*(1-K)))
        return None

    def _clamp(c): return tuple(max(0,min(255,int(v))) for v in c) if c else None

    def _resolve_font(name, bold, italic):
        for kw in ("-BoldItalic","BoldItalic"):
            if kw in name: name=name.replace(kw,""); bold=italic=True
        for kw in ("-Bold","Bold"):
            if kw in name: name=name.replace(kw,""); bold=True
        for kw in ("-Italic","Italic","-Oblique","Oblique"):
            if kw in name: name=name.replace(kw,""); italic=True
        name=name.strip("-_ ")
        for k,v in JP_FONT_MAP.items():
            if k in name.lower(): return (v,bold,italic)
        return (name or "Arial",bold,italic)

    with tempfile.NamedTemporaryFile(suffix=".pdf",delete=False) as tmp:
        tmp.write(pdf_bytes); tmp_path=tmp.name

    try:
        doc=fitz.open(tmp_path); n=len(doc)
        first=doc[0]; aspect=first.rect.height/first.rect.width; doc.close()

        prs=Presentation()
        prs.slide_width=Inches(slide_width_inches)
        prs.slide_height=int(Inches(slide_width_inches)*aspect)
        sw,sh=int(prs.slide_width),int(prs.slide_height)
        stats={"pages":n,"texts":0,"images":0,"shapes":0,"is_image_pdf":False}

        for i in range(n):
            if progress_cb: progress_cb(i,n,f"ページ {i+1}/{n} を変換中...")
            doc2=fitz.open(tmp_path); page=doc2[i]
            pw,ph=page.rect.width,page.rect.height

            # テキストが0で画像が1つ = 画像化PDFと判定
            raw_dict=page.get_text("dict",flags=fitz.TEXT_PRESERVE_WHITESPACE)
            text_blocks=[b for b in raw_dict.get("blocks",[]) if b.get("type")==0]
            imgs=page.get_images(full=True)
            is_image_page = len(text_blocks)==0 and len(imgs)>=1

            if is_image_page:
                stats["is_image_pdf"] = True

            def ex(v): return int(v*PT_TO_EMU*(sw/int(pw*PT_TO_EMU)))
            def ey(v): return int(v*PT_TO_EMU*(sh/int(ph*PT_TO_EMU)))
            def ew(v): return max(1,int(v*PT_TO_EMU*(sw/int(pw*PT_TO_EMU))))
            def eh(v): return max(1,int(v*PT_TO_EMU*(sh/int(ph*PT_TO_EMU))))
            def clamp(l,t,w,h):
                l=max(0,min(l,sw-1)); t=max(0,min(t,sh-1))
                return l,t,max(1,min(w,sw-l)),max(1,min(h,sh-t))

            slide=prs.slides.add_slide(prs.slide_layouts[6])

            if is_image_page:
                # 画像化ページ：高画質でラスタライズして配置
                pix=page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False)
                img_bytes=pix.tobytes("png")
                slide.shapes.add_picture(io.BytesIO(img_bytes),0,0,sw,sh)
                stats["images"]+=1
                doc2.close()
                continue

            # 背景色
            for path in page.get_drawings():
                r=path.get("rect")
                if r and r.x0<=2 and r.y0<=2 and r.x1>=pw-2 and r.y1>=ph-2:
                    fill=_clamp(_to_rgb(path.get("fill")))
                    if fill:
                        rf,gf,bf=fill
                        NP="http://schemas.openxmlformats.org/presentationml/2006/main"
                        NA="http://schemas.openxmlformats.org/drawingml/2006/main"
                        xml=(f'<p:bg xmlns:p="{NP}" xmlns:a="{NA}"><p:bgPr>'
                             f'<a:solidFill><a:srgbClr val="{rf:02X}{gf:02X}{bf:02X}"/></a:solidFill>'
                             f'<a:effectLst/></p:bgPr></p:bg>')
                        slide.shapes._spTree.getparent().insert(2,etree.fromstring(xml))
                        break

            # 図形
            for path in page.get_drawings():
                r=path.get("rect")
                if not r: continue
                x0,y0=max(0.,float(r.x0)),max(0.,float(r.y0))
                x1,y1=min(pw,float(r.x1)),min(ph,float(r.y1))
                if x1-x0<3 or y1-y0<3: continue
                if x0<=2 and y0<=2 and x1>=pw-2 and y1>=ph-2: continue
                fill=_clamp(_to_rgb(path.get("fill")))
                stroke=_clamp(_to_rgb(path.get("color")))
                if not fill and not stroke: continue
                l,t,w,h=clamp(ex(x0),ey(y0),ew(x1-x0),eh(y1-y0))
                try:
                    shp=slide.shapes.add_shape(1,l,t,w,h)
                    if fill: shp.fill.solid(); shp.fill.fore_color.rgb=RGBColor(*fill)
                    else: shp.fill.background()
                    if stroke: shp.line.color.rgb=RGBColor(*stroke); shp.line.width=Pt(max(float(path.get("width") or 1),.25))
                    else: shp.line.fill.background()
                    stats["shapes"]+=1
                except: pass

            # 画像
            for info in page.get_images(full=True):
                xref=info[0]; rects=page.get_image_rects(xref)
                if not rects: continue
                r=rects[0]
                if (r.x1-r.x0)<5 or (r.y1-r.y0)<5: continue
                try:
                    d=doc2.extract_image(xref)
                    raw,ext=d.get("image",b""),d.get("ext","png").lower()
                    if not raw: continue
                    if ext not in ("png","jpeg","jpg"):
                        buf=io.BytesIO(); PILImage.open(io.BytesIO(raw)).save(buf,"PNG"); raw=buf.getvalue()
                    l,t,w,h=clamp(ex(r.x0),ey(r.y0),ew(r.x1-r.x0),eh(r.y1-r.y0))
                    slide.shapes.add_picture(io.BytesIO(raw),l,t,w,h)
                    stats["images"]+=1
                except: pass

            # テキスト
            sx_scale=sw/int(pw*PT_TO_EMU); sy_scale=sh/int(ph*PT_TO_EMU)
            for blk in text_blocks:
                bbox=blk["bbox"]; bx=(bbox[0]+bbox[2])/2
                spans_out,aligns=[],[]
                for line in blk.get("lines",[]):
                    lx=(line["bbox"][0]+line["bbox"][2])/2; d=lx-bx
                    aligns.append("center" if abs(d)<15 else ("right" if d>25 else "left"))
                    for span in line.get("spans",[]):
                        txt=span.get("text","")
                        if not txt: continue
                        fl=span.get("flags",0)
                        c=_clamp(_to_rgb(span.get("color",0)))
                        spans_out.append((txt,span.get("font","Arial"),span.get("size",12),
                                        bool(fl&(1<<4)),bool(fl&(1<<1)),c or (0,0,0)))
                if not spans_out: continue
                align=Counter(aligns).most_common(1)[0][0] if aligns else "left"
                l,t,w,h=clamp(ex(bbox[0]),ey(bbox[1]),ew(max(bbox[2]-bbox[0],20)+20),eh(max(bbox[3]-bbox[1],10)+8))
                txb=slide.shapes.add_textbox(l,t,w,h)
                tf=txb.text_frame; tf.word_wrap=False
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=Pt(0)
                AMAP={"left":PP_ALIGN.LEFT,"center":PP_ALIGN.CENTER,"right":PP_ALIGN.RIGHT}
                para=tf.paragraphs[0]; para.alignment=AMAP.get(align,PP_ALIGN.LEFT)
                fs_scale=min(sx_scale,sy_scale)
                for txt,fname,fsize,bold,italic,color in spans_out:
                    for li,ln in enumerate(txt.split("\n")):
                        if li>0: para=tf.add_paragraph(); para.alignment=AMAP.get(align,PP_ALIGN.LEFT)
                        if not ln: continue
                        run=para.add_run(); run.text=ln
                        fn,bd,it=_resolve_font(fname,bold,italic)
                        f=run.font; f.name=fn; f.size=Pt(fsize*fs_scale); f.bold=bd; f.italic=it
                        f.color.rgb=RGBColor(*color)
                stats["texts"]+=1
            doc2.close()

        buf=io.BytesIO(); prs.save(buf)
        return buf.getvalue(), stats
    finally:
        os.unlink(tmp_path)


# ── 変換エンジン（ハード・AI） ───────────────────────
def analyze_page_with_gemini(img_bytes, api_key):
    import google.generativeai as genai
    genai.configure(api_key=api_key)
    model = genai.GenerativeModel("gemini-2.0-flash")
    prompt = """このPowerPointスライドの画像を詳細に解析してください。

以下のJSON形式のみで返してください（説明文・マークダウン不要）:

{
  "background_color": "#RRGGBB",
  "elements": [
    {
      "type": "text",
      "text": "テキスト内容（改行は\\nで）",
      "x": 0.05, "y": 0.05, "w": 0.9, "h": 0.1,
      "font_size": 24, "bold": true,
      "color": "#ffffff", "bg_color": null,
      "border_color": null, "align": "center"
    },
    {
      "type": "rect",
      "x": 0.0, "y": 0.8, "w": 1.0, "h": 0.2,
      "fill_color": "#003366", "border_color": null
    },
    {
      "type": "image",
      "x": 0.1, "y": 0.2, "w": 0.4, "h": 0.5,
      "description": "天秤の写真"
    }
  ]
}

ルール:
- x,y,w,h は画像全体を1.0とした相対値（左上が原点）
- すべてのテキストを漏れなく抽出（画像に焼き付いた文字も含む）
- 背景の色付きボックスや図形も rect として含める
- 写真・イラスト・アイコンは image として含める
- font_sizeは実際のサイズを推定
- JSONのみ返すこと"""

    response = model.generate_content([
        {"mime_type": "image/png", "data": base64.b64encode(img_bytes).decode()},
        prompt
    ])
    raw = response.text.strip()
    raw = re.sub(r"^```[a-zA-Z]*\n?", "", raw)
    raw = re.sub(r"\n?```$", "", raw)
    raw = raw.strip()
    try:
        return json.loads(raw)
    except:
        m = re.search(r'\{.*\}', raw, re.DOTALL)
        if m: return json.loads(m.group())
        raise


def convert_hard(pdf_bytes, api_key, slide_width_inches=13.333, progress_cb=None):
    import fitz
    from pptx import Presentation
    from pptx.util import Pt, Inches
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from lxml import etree

    def h2r(h):
        if not h or len(h.lstrip("#"))!=6: return (0,0,0)
        h=h.lstrip("#")
        return tuple(int(h[i:i+2],16) for i in (0,2,4))

    with tempfile.NamedTemporaryFile(suffix=".pdf",delete=False) as tmp:
        tmp.write(pdf_bytes); tmp_path=tmp.name

    try:
        doc=fitz.open(tmp_path); n=len(doc)
        first=doc[0]; aspect=first.rect.height/first.rect.width; doc.close()

        prs=Presentation()
        prs.slide_width=Inches(slide_width_inches)
        prs.slide_height=int(Inches(slide_width_inches)*aspect)
        sw,sh=int(prs.slide_width),int(prs.slide_height)
        stats={"pages":n,"texts":0,"images":0,"shapes":0,"errors":[]}

        for i in range(n):
            if progress_cb: progress_cb(i,n,f"ページ {i+1}/{n} を画像化中...")
            doc2=fitz.open(tmp_path)
            pix=doc2[i].get_pixmap(matrix=fitz.Matrix(2,2),alpha=False)
            img_bytes=pix.tobytes("png")
            page_w=doc2[i].rect.width; page_h=doc2[i].rect.height
            doc2.close()

            if progress_cb: progress_cb(i,n,f"ページ {i+1}/{n} をAIが解析中... ⏳")
            try:
                analysis=analyze_page_with_gemini(img_bytes,api_key)
            except Exception as e:
                stats["errors"].append(f"P{i+1}: {str(e)[:60]}")
                slide=prs.slides.add_slide(prs.slide_layouts[6])
                slide.shapes.add_picture(io.BytesIO(img_bytes),0,0,sw,sh)
                continue

            slide=prs.slides.add_slide(prs.slide_layouts[6])

            def rx(v): return max(0,int(float(v)*sw))
            def ry(v): return max(0,int(float(v)*sh))
            def clamp(l,t,w,h):
                l=max(0,min(int(l),sw-1)); t=max(0,min(int(t),sh-1))
                return l,t,max(1,min(int(w),sw-l)),max(1,min(int(h),sh-t))

            bg=analysis.get("background_color","")
            if bg and len(bg)==7:
                try:
                    r,g,b=h2r(bg)
                    NP="http://schemas.openxmlformats.org/presentationml/2006/main"
                    NA="http://schemas.openxmlformats.org/drawingml/2006/main"
                    xml=(f'<p:bg xmlns:p="{NP}" xmlns:a="{NA}"><p:bgPr>'
                         f'<a:solidFill><a:srgbClr val="{r:02X}{g:02X}{b:02X}"/></a:solidFill>'
                         f'<a:effectLst/></p:bgPr></p:bg>')
                    slide.shapes._spTree.getparent().insert(2,etree.fromstring(xml))
                except: pass

            AMAP={"left":PP_ALIGN.LEFT,"center":PP_ALIGN.CENTER,"right":PP_ALIGN.RIGHT}
            elements=analysis.get("elements",[])
            if progress_cb: progress_cb(i,n,f"ページ {i+1}/{n}: {len(elements)}個の要素を配置中...")

            for el in elements:
                try:
                    et=el.get("type","")
                    l,t,w,h=clamp(rx(el.get("x",0)),ry(el.get("y",0)),
                                  rx(el.get("w",0.1)),ry(el.get("h",0.1)))
                    if et=="rect":
                        shp=slide.shapes.add_shape(1,l,t,w,h)
                        fc=el.get("fill_color")
                        if fc: shp.fill.solid(); shp.fill.fore_color.rgb=RGBColor(*h2r(fc))
                        else: shp.fill.background()
                        bc=el.get("border_color")
                        if bc: shp.line.color.rgb=RGBColor(*h2r(bc)); shp.line.width=Pt(1)
                        else: shp.line.fill.background()
                        stats["shapes"]+=1
                    elif et=="image":
                        doc3=fitz.open(tmp_path)
                        page=doc3[i]
                        clip=fitz.Rect(
                            el.get("x",0)*page_w, el.get("y",0)*page_h,
                            (el.get("x",0)+el.get("w",0.1))*page_w,
                            (el.get("y",0)+el.get("h",0.1))*page_h
                        )
                        pix2=page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False,clip=clip)
                        region_bytes=pix2.tobytes("png"); doc3.close()
                        slide.shapes.add_picture(io.BytesIO(region_bytes),l,t,w,h)
                        stats["images"]+=1
                    elif et=="text":
                        txt=el.get("text","")
                        if not txt: continue
                        txb=slide.shapes.add_textbox(l,t,w,h)
                        tf=txb.text_frame; tf.word_wrap=True
                        tf.margin_left=tf.margin_right=Pt(3)
                        tf.margin_top=tf.margin_bottom=Pt(2)
                        bg_c=el.get("bg_color")
                        if bg_c: txb.fill.solid(); txb.fill.fore_color.rgb=RGBColor(*h2r(bg_c))
                        bc=el.get("border_color")
                        if bc: txb.line.color.rgb=RGBColor(*h2r(bc)); txb.line.width=Pt(1.5)
                        align=AMAP.get(el.get("align","left"),PP_ALIGN.LEFT)
                        fs=max(float(el.get("font_size",16))*0.72,8)
                        lines=txt.split("\n")
                        para=tf.paragraphs[0]; para.alignment=align
                        for li,line in enumerate(lines):
                            if li>0: para=tf.add_paragraph(); para.alignment=align
                            if not line.strip(): continue
                            run=para.add_run(); run.text=line
                            f=run.font; f.size=Pt(fs); f.bold=el.get("bold",False)
                            f.color.rgb=RGBColor(*h2r(el.get("color","#000000")))
                        stats["texts"]+=1
                except: pass

        buf=io.BytesIO(); prs.save(buf)
        return buf.getvalue(), stats
    finally:
        os.unlink(tmp_path)


# ── 変換実行ボタン ───────────────────────────────────
if uploaded:
    step_num2 = "④" if is_hard else "③"
    st.markdown(f'<div class="section-card"><h3>ステップ{step_num2} 変換を実行</h3>', unsafe_allow_html=True)

    can_convert = True
    if is_hard and not gemini_key:
        st.warning("⬆️ APIキーを入力してから変換してください")
        can_convert = False

    if can_convert:
        btn_label = "🤖　AIモードで変換開始" if is_hard else "⚡　ライトモードで変換開始"
        if st.button(btn_label, use_container_width=True):
            pdf_bytes = uploaded.read()
            progress_bar = st.progress(0.0)
            status_text = st.empty()

            import fitz as _fitz
            _doc=_fitz.open(stream=pdf_bytes,filetype="pdf")
            n_pages=len(_doc); _doc.close()

            def progress_cb(page_i, total, msg):
                progress_bar.progress(min((page_i+0.5)/total, 0.99))
                status_text.markdown(f"**{msg}**")

            try:
                if is_hard:
                    pptx_bytes, stats = convert_hard(pdf_bytes, gemini_key, progress_cb=progress_cb)
                else:
                    pptx_bytes, stats = convert_light(pdf_bytes, progress_cb=progress_cb)

                progress_bar.progress(1.0)
                status_text.markdown("**✅ 変換完了！**")

                # 画像化PDFの警告（ライトモード時）
                if not is_hard and stats.get("is_image_pdf"):
                    st.warning("⚠️ このPDFは画像化されています。テキストを編集可能にするには**ハードモード（AI）**をお試しください。")

                if stats.get("errors"):
                    st.warning("一部エラー：" + " / ".join(stats["errors"]))

                st.markdown(f"""
<div class="section-card">
<div style="font-size:1rem;font-weight:700;color:#333;margin-bottom:0.8rem;">✨ 変換完了！</div>
<div class="stat-grid">
  <div class="stat-box"><div class="stat-num">{stats['pages']}</div><div class="stat-lbl">スライド</div></div>
  <div class="stat-box"><div class="stat-num">{stats['texts']}</div><div class="stat-lbl">編集可能テキスト</div></div>
  <div class="stat-box"><div class="stat-num">{stats['images']}</div><div class="stat-lbl">画像</div></div>
</div>
</div>
""", unsafe_allow_html=True)

                out_name = uploaded.name.replace(".pdf","_変換済み.pptx").replace(".PDF","_変換済み.pptx")
                st.download_button(
                    label="⬇️　PPTXをダウンロード",
                    data=pptx_bytes,
                    file_name=out_name,
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    use_container_width=True,
                )

            except Exception as e:
                import traceback
                progress_bar.empty(); status_text.empty()
                st.error(f"❌ エラー: {e}")
                with st.expander("エラー詳細"):
                    st.code(traceback.format_exc())

    st.markdown('</div>', unsafe_allow_html=True)

st.markdown("---")
st.markdown("<p style='text-align:center;color:#aaa;font-size:0.78rem;'>PDF → PPTX Converter v4.1 | ⚡ライト＆🤖ハードモード対応</p>", unsafe_allow_html=True)
